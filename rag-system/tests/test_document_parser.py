import fitz
import pytest
from docx import Document

from app.parser.document_parser import _normalize_block, parse_document


def _make_pdf(path, pages):
    """Write a PDF where each page is (header, body, footer); None skips a part."""
    pdf = fitz.open()
    for header, body, footer in pages:
        page = pdf.new_page()  # A4: 595 x 842
        if header:
            page.insert_text((72, 30), header, fontsize=10)
        if body:
            page.insert_text((72, 400), body, fontsize=11)
        if footer:
            page.insert_text((72, 820), footer, fontsize=10)
    pdf.save(path)
    pdf.close()


def test_normalize_block_ignores_page_numbers_and_whitespace():
    assert _normalize_block("Page 3 of 6\n") == _normalize_block("Page  4 of 6")


def test_parse_txt(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Personal data is processed.", encoding="utf-8")

    result = parse_document(str(path))

    assert result["text"] == "Personal data is processed."
    assert result["metadata"] == {"source": str(path), "file_type": ".txt"}


def test_parse_docx(tmp_path):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("Clause 1: Data retention")
    doc.add_paragraph("Clause 2: Subprocessors")
    doc.save(path)

    result = parse_document(str(path))

    assert result["text"] == "Clause 1: Data retention\nClause 2: Subprocessors"
    assert result["metadata"]["file_type"] == ".docx"


def test_extension_is_case_insensitive(tmp_path):
    path = tmp_path / "NOTES.TXT"
    path.write_text("hello", encoding="utf-8")

    assert parse_document(str(path))["metadata"]["file_type"] == ".txt"


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "sheet.xlsx"
    path.write_bytes(b"")

    with pytest.raises(ValueError, match="Unsupported file type: .xlsx"):
        parse_document(str(path))


def test_pdf_strips_repeated_headers_and_footers(tmp_path):
    path = tmp_path / "dpa.pdf"
    _make_pdf(path, [
        ("ACME Confidential", f"Body clause {i}", f"Page {i} of 3")
        for i in range(1, 4)
    ])

    text = parse_document(str(path))["text"]

    for i in range(1, 4):
        assert f"Body clause {i}" in text
    assert "ACME Confidential" not in text
    assert "Page" not in text


def test_pdf_keeps_margin_text_that_does_not_repeat(tmp_path):
    path = tmp_path / "dpa.pdf"
    _make_pdf(path, [
        ("Schedule A", "First body", None),
        ("Schedule B", "Second body", None),
        ("Schedule C", "Third body", None),
    ])

    text = parse_document(str(path))["text"]

    for label in ("Schedule A", "Schedule B", "Schedule C"):
        assert label in text


def test_single_page_pdf_keeps_margin_text(tmp_path):
    path = tmp_path / "one.pdf"
    _make_pdf(path, [("ACME Confidential", "Only body", "Page 1 of 1")])

    text = parse_document(str(path))["text"]

    assert "ACME Confidential" in text
    assert "Only body" in text
    assert "Page 1 of 1" in text


def test_pdf_blocks_are_newline_separated(tmp_path):
    path = tmp_path / "one.pdf"
    _make_pdf(path, [(None, "Only body", None)])

    assert parse_document(str(path))["text"].endswith("Only body\n")
